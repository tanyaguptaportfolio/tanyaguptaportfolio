from docx import Document

from backend.extract_cv import extract_document


def test_extracts_paragraphs_and_tables(tmp_path):
    path = tmp_path / "profile.docx"
    document = Document()
    document.add_paragraph("Tanya")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Verified role"
    document.sections[0].header.paragraphs[0].text = "CV header"
    document.save(path)
    content = extract_document(path)
    assert content["source"] == "profile.docx"
    assert content["paragraphs"] == ["Tanya", "Role | Verified role", "CV header"]
